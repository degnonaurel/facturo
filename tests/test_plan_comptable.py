# -*- coding: utf-8 -*-
"""Tests du plan comptable : lecture, catégorisation, onglet et liste déroulante."""
from _helpers import RACINE, pdf_facture_demo  # noqa: F401
import facture_vers_excel as fve
from openpyxl import Workbook, load_workbook


def test_plan_par_defaut_bilingue():
    fr, en = fve.plan_par_defaut("fr"), fve.plan_par_defaut("en")
    assert len(fr) == len(en) and fr[0]["code"] == en[0]["code"]
    location = [c for c in fr if c["code"] == "5320"][0]
    assert location["nom"] == "Location de véhicule" and location["type"] == "Dépenses"
    assert [c for c in en if c["code"] == "5320"][0]["nom"] == "Vehicle rental"


def test_lire_plan_csv_avec_titres(tmp_path):
    p = tmp_path / "plan.csv"
    p.write_text("No compte;Libellé;Classe\n6100;Repas d'affaires;Charges\n"
                 "6200;Location auto;Charges\n", encoding="utf-8")
    plan = fve.lire_plan_comptable(str(p))
    assert plan == [{"code": "6100", "nom": "Repas d'affaires", "type": "Charges"},
                    {"code": "6200", "nom": "Location auto", "type": "Charges"}]


def test_lire_plan_excel_sans_titres(tmp_path):
    p = str(tmp_path / "plan.xlsx")
    wb = Workbook()
    for code, nom in ((6100, "Repas"), (6200, "Location auto")):
        wb.active.append([code, nom])
    wb.save(p)
    assert [c["nom"] for c in fve.lire_plan_comptable(p)] == ["Repas", "Location auto"]
    assert fve.lire_plan_comptable(p)[1]["code"] == "6200"


def test_plan_liste_de_noms(tmp_path):
    p = tmp_path / "cats.csv"
    p.write_text("Repas\nLocation auto\n", encoding="utf-8")
    plan = fve.lire_plan_comptable(str(p))
    assert plan[1] == {"code": "", "nom": "Location auto", "type": ""}
    assert fve.libelle_compte(plan, "Location auto") == "Location auto"


def test_schema_envoye_contient_les_comptes(monkeypatch):
    vu = {}

    def faux_claude(texte, images, modele, timeout, schema, consigne):
        vu["schema"], vu["consigne"] = schema, consigne
        return {}
    monkeypatch.setenv("ANTHROPIC_API_KEY", "cle-test")
    monkeypatch.setattr(fve, "_appeler_claude", faux_claude)
    plan = fve.plan_par_defaut("fr")
    fve.extraire_avec_llm(texte="x", plan=plan)
    assert "5320" in vu["schema"]["properties"]["compte"]["enum"]
    assert "compte" in vu["schema"]["required"]
    assert "Location de véhicule" in vu["consigne"]
    assert "compte" not in fve._SCHEMA_FACTURE["properties"]   # schéma de base intact


def test_categorisation_bout_en_bout(tmp_path, monkeypatch):
    """Location auto → « 5320 · Location de véhicule » dans la colonne
    Catégorie, onglet plan + liste déroulante, plan réutilisé à l'ajout."""
    pdf = pdf_facture_demo(str(tmp_path / "f.pdf"))
    faux = {"fournisseur": "Enterprise", "numero": "R-1", "date": "2026-09-01",
            "devise": "CAD", "total_ht": 100.0, "tps": 5.0, "tvq": 9.98,
            "tva": 14.98, "total_ttc": 114.98, "lignes": [], "compte": "5320"}
    monkeypatch.setenv("ANTHROPIC_API_KEY", "cle-test")
    monkeypatch.setattr(fve, "extraire_avec_llm", lambda *a, **k: faux)
    plan = fve.plan_par_defaut("fr")
    f = fve.traiter_facture(pdf, moteur="llm", plan=plan)
    assert f.categorie == "5320 · Location de véhicule"

    base = str(tmp_path / "compta.xlsx")
    fve.construire_excel([f], base, plan=plan)
    wb = load_workbook(base)
    assert wb.sheetnames == ["Resume", "Details", "Totaux par catégorie", "Plan comptable"]
    ws = wb["Resume"]
    col = [c.value for c in ws[1]].index("Catégorie") + 1
    assert ws.cell(2, col).value == "5320 · Location de véhicule"
    dv = ws.data_validations.dataValidation
    assert len(dv) == 1 and "'Plan comptable'!$D$2" in dv[0].formula1

    # Ajout en anglais sans fournir de plan : celui du classeur est repris.
    assert [c["code"] for c in fve.plan_du_classeur(base)] == [c["code"] for c in plan]
    fve.construire_excel([f], base, base_excel=base, langue="en",
                         plan=fve.plan_du_classeur(base))
    wb = load_workbook(base)
    assert wb.sheetnames == ["Summary", "Details", "Totals by category", "Chart of accounts"]
    dv = wb["Summary"].data_validations.dataValidation
    assert len(dv) == 1 and "'Chart of accounts'!" in dv[0].formula1


def _piece(nom, numero, ht, tx, ttc, cat, lignes):
    plan = fve.plan_par_defaut("fr")
    brut = {"fournisseur": nom, "numero": numero, "date": "2026-09-01",
            "devise": "CAD", "total_ht": ht, "tps": None, "tvq": None, "tva": tx,
            "total_ttc": ttc, "compte": cat,
            "lignes": [{"description": d, "montant": m, "compte": c}
                       for d, m, c in lignes]}
    return fve._vers_facture(nom + ".pdf", brut, "test", plan), plan


def test_categorie_par_article_et_schema(monkeypatch):
    f, _ = _piece("Costco", "C1", 100, 13, 113, "5240",
                  [("Papier", 75, "5240"), ("Sandwich", 25, "5210")])
    assert [l.categorie for l in f.lignes] == ["5240 · Fournitures de bureau",
                                               "5210 · Repas et représentation"]
    vu = {}
    monkeypatch.setenv("ANTHROPIC_API_KEY", "cle-test")
    monkeypatch.setattr(fve, "_appeler_claude",
                        lambda t, i, m, to, schema, consigne: vu.update(s=schema) or {})
    fve.extraire_avec_llm(texte="x", plan=fve.plan_par_defaut("fr"))
    article = vu["s"]["properties"]["lignes"]["items"]
    assert "compte" in article["required"] and "5210" in article["properties"]["compte"]["enum"]


def test_totaux_par_categorie_avec_repartition(tmp_path):
    """Reçu mixte réparti au prorata des articles ; pièce simple entière ;
    onglet recalculé à chaque ajout ; TOTAL identique au résumé."""
    mixte, plan = _piece("Costco", "C1", 100, 13, 113, "5240",
                         [("Papier", 75, "5240"), ("Sandwich", 25, "5210")])
    simple, _ = _piece("Enterprise", "E1", 200, 26, 226, "5320",
                       [("Location 2 jours", 200, "5320")])
    base = str(tmp_path / "compta.xlsx")
    fve.construire_excel([mixte], base, plan=plan)
    fve.construire_excel([simple], base, base_excel=base, plan=plan)

    wt = load_workbook(base)["Totaux par catégorie"]
    lignes = {r[0]: r for r in wt.iter_rows(min_row=2, values_only=True)}
    assert lignes["5240 · Fournitures de bureau"][2:] == (1, 75, 0, 0, 9.75, 84.75)
    assert lignes["5210 · Repas et représentation"][2:] == (1, 25, 0, 0, 3.25, 28.25)
    assert lignes["5320 · Location de véhicule"][2:] == (1, 200, 0, 0, 26, 226)
    assert lignes["TOTAL CAD"][2:] == (2, 300, 0, 0, 39, 339)
    assert len(lignes) == 4                      # pas de doublon après l'ajout


def test_traitement_en_parallele_garde_l_ordre(monkeypatch):
    import threading
    import time
    actifs, maxi = [0], [0]
    verrou = threading.Lock()

    def lent(chemin, **options):
        with verrou:
            actifs[0] += 1
            maxi[0] = max(maxi[0], actifs[0])
        time.sleep(0.05)
        with verrou:
            actifs[0] -= 1
        if chemin == "mauvais":
            raise fve.ErreurLLM("illisible")
        return chemin.upper()
    monkeypatch.setattr(fve, "traiter_facture", lent)
    res = fve.traiter_lot(["a", "mauvais", "c", "d"], paralleles=4)
    assert res[0] == "A" and isinstance(res[1], fve.ErreurLLM) and res[2:] == ["C", "D"]
    assert maxi[0] > 1                           # vraiment en parallèle
