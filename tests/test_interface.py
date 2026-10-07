# -*- coding: utf-8 -*-
"""Tests de l'app web (TestClient FastAPI, sans réseau ni clé)."""
import io
from _helpers import RACINE, pdf_facture_demo  # noqa: F401
from fastapi.testclient import TestClient
import interface_web
from openpyxl import load_workbook

client = TestClient(interface_web.app)


def test_page_sans_mention_moteur():
    page = client.get("/").text
    assert "Facturo" in page
    assert "Claude" not in page and "Anthropic" not in page
    assert "data-i18n" in page  # bilingue


def test_statut():
    assert "pret" in client.get("/api/statut").json()


def test_extraire_et_telecharger(tmp_path):
    pdf = pdf_facture_demo(str(tmp_path / "f.pdf"))
    with open(pdf, "rb") as fh:
        files = [("fichiers", ("f.pdf", fh.read(), "application/pdf"))]
    d = client.post("/api/extraire", files=files).json()
    assert len(d["factures"]) == 1 and d["download_id"]
    xl = client.get("/telecharger/" + d["download_id"]).content
    wb = load_workbook(io.BytesIO(xl))
    assert wb.sheetnames == ["Resume", "Details", "Totaux par catégorie", "Plan comptable"]


def test_ajout_a_base_via_web(tmp_path):
    # 1) produire un premier Excel
    import facture_vers_excel as fve
    base = str(tmp_path / "base.xlsx")
    d0 = {"fournisseur": "Alpha", "numero": "1", "date": "2026-09-01",
          "devise": "CAD", "total_ht": None, "tps": None, "tvq": None,
          "tva": None, "total_ttc": 10.0, "lignes": [{"description": "x", "montant": 10}]}
    fve.construire_excel([fve._vers_facture("a.pdf", d0, "t")], base)

    pdf = pdf_facture_demo(str(tmp_path / "f.pdf"))
    with open(pdf, "rb") as fh, open(base, "rb") as bh:
        files = [("fichiers", ("f.pdf", fh.read(), "application/pdf")),
                 ("base", ("base.xlsx", bh.read(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))]
    d = client.post("/api/extraire", files=files).json()
    assert d["ajoute"] is True
    xl = client.get("/telecharger/" + d["download_id"]).content
    wb = load_workbook(io.BytesIO(xl))
    pieces = [r for r in wb["Resume"].iter_rows(min_row=2, values_only=True)
              if not fve._est_ligne_total(r[0])]
    assert len(pieces) == 2  # 1 existante + 1 nouvelle (+ lignes TOTAL)


def test_quota_de_demo(tmp_path, monkeypatch):
    monkeypatch.setattr(interface_web, "QUOTA_JOUR_VISITEUR", 2)
    monkeypatch.setattr(interface_web, "_quota",
                        {"jour": None, "total": 0, "visiteurs": {}})
    pdf = pdf_facture_demo(str(tmp_path / "f.pdf"))
    with open(pdf, "rb") as fh:
        contenu = fh.read()

    def envoyer(n, ip):
        files = [("fichiers", (f"f{i}.pdf", contenu, "application/pdf"))
                 for i in range(n)]
        return client.post("/api/extraire", files=files,
                           headers={"X-Forwarded-For": ip})

    assert envoyer(2, "1.1.1.1").status_code == 200
    r = envoyer(1, "1.1.1.1")                    # 3e fichier du jour : refusé
    assert r.status_code == 429 and r.json()["quota_atteint"] is True
    assert envoyer(1, "2.2.2.2").status_code == 200   # autre visiteur : OK


def test_langue_anglaise_via_web(tmp_path):
    pdf = pdf_facture_demo(str(tmp_path / "f.pdf"))
    with open(pdf, "rb") as fh:
        files = [("fichiers", ("f.pdf", fh.read(), "application/pdf"))]
    d = client.post("/api/extraire", files=files, data={"langue": "en"}).json()
    wb = load_workbook(io.BytesIO(client.get("/telecharger/" + d["download_id"]).content))
    assert wb.sheetnames == ["Summary", "Details", "Totals by category", "Chart of accounts"]
    assert wb["Summary"]["A1"].value == "File"


def test_totaux_fichier_renvoyes(tmp_path):
    pdf = pdf_facture_demo(str(tmp_path / "f.pdf"))
    with open(pdf, "rb") as fh:
        files = [("fichiers", ("f.pdf", fh.read(), "application/pdf"))]
    d = client.post("/api/extraire", files=files).json()
    assert d["totaux_fichier"] and d["totaux_fichier"][0]["nb"] == 1


def test_statut_quota_et_contact(tmp_path, monkeypatch):
    monkeypatch.setattr(interface_web, "_quota", {"jour": None, "total": 0, "visiteurs": {}})
    s = client.get("/api/statut", headers={"X-Forwarded-For": "9.9.9.9"}).json()
    assert s["quota"] == interface_web.QUOTA_JOUR_VISITEUR
    assert s["restant"] == interface_web.QUOTA_JOUR_VISITEUR and "@" in s["contact"]
    pdf = pdf_facture_demo(str(tmp_path / "f.pdf"))
    with open(pdf, "rb") as fh:
        files = [("fichiers", ("f.pdf", fh.read(), "application/pdf"))]
    d = client.post("/api/extraire", files=files, headers={"X-Forwarded-For": "9.9.9.9"}).json()
    assert d["restant"] == interface_web.QUOTA_JOUR_VISITEUR - 1
    page = client.get("/").text
    assert 'id="offre"' in page and "Facturo Pro" in page


def test_dockerfile_copie_tous_les_modules():
    """L'image de déploiement contient chaque module Python de la racine
    (sinon la démo plante au démarrage)."""
    import re
    copie = " ".join(re.findall(r"^COPY (.+)$", (RACINE / "Dockerfile").read_text(), re.M))
    for module in RACINE.glob("*.py"):
        assert module.name in copie, f"{module.name} absent du Dockerfile"
