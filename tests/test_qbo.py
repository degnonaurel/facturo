# -*- coding: utf-8 -*-
"""Tests de l'export CSV QuickBooks Online (import de factures fournisseurs)."""
import csv
import io
from _helpers import RACINE, pdf_facture_demo  # noqa: F401
import facture_vers_excel as fve


def _facture(**champs):
    brut = {"fournisseur": "Bureau en Gros", "numero": "A-1", "date": "2026-09-03",
            "devise": "CAD", "total_ht": 100.0, "tps": 5.0, "tvq": 9.98,
            "tva": 14.98, "total_ttc": 114.98, "lignes": []}
    brut.update(champs)
    return fve._vers_facture("recu.jpg", brut, "t", fve.plan_par_defaut("fr"))


def _lire(chemin):
    with open(chemin, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def test_codes_de_taxe():
    assert fve.code_taxe_qbo(_facture()) == "GST/QST QC - 9.975"
    assert fve.code_taxe_qbo(_facture(tvq=None, tva=5.0, total_ttc=105.0)) == "GST"
    assert fve.code_taxe_qbo(_facture(tps=None, tvq=None, tva=13.0,
                                      total_ttc=113.0)) == "HST ON"
    assert fve.code_taxe_qbo(_facture(tps=None, tvq=None, tva=0,
                                      total_ttc=100.0)) == "Exempt"
    assert fve.code_taxe_qbo(_facture(devise="USD", tps=None, tvq=None, tva=None,
                                      total_ttc=100.0)) == "Out of Scope"


def test_csv_une_ligne_par_article(tmp_path):
    f = _facture(lignes=[
        {"description": "Papier", "montant": 60.0, "compte": "5240"},
        {"description": "Lunch", "montant": 40.0, "compte": "5210"}])
    chemin = fve.construire_csv_qbo([f], str(tmp_path / "qbo.csv"),
                                    plan=fve.plan_par_defaut("fr"))
    lignes = _lire(chemin)
    assert list(lignes[0].keys()) == fve.COLS_QBO
    assert [l["Account"] for l in lignes] == ["Fournitures de bureau",
                                              "Repas et représentation"]
    assert all(l["Bill no."] == "A-1" and l["Supplier"] == "Bureau en Gros"
               and l["Bill Date"] == "03/09/2026" and l["Due Date"] == "03/09/2026"
               for l in lignes)
    assert [l["Line Amount"] for l in lignes] == ["60.00", "40.00"]
    # Taxes réparties au prorata, sans perte d'arrondi.
    assert round(sum(float(l["Line Tax Amount"]) for l in lignes), 2) == 14.98


def test_csv_une_ligne_si_articles_incoherents(tmp_path):
    f = _facture(numero="", lignes=[{"description": "x", "montant": 3.0}])
    lignes = _lire(fve.construire_csv_qbo([f], str(tmp_path / "q.csv")))
    assert len(lignes) == 1 and lignes[0]["Line Amount"] == "100.00"
    assert lignes[0]["Bill no."] == "recu"          # n° absent → nom du fichier
    assert lignes[0]["Line Tax Code"] == "GST/QST QC - 9.975"


def test_csv_via_web_seulement_sur_demande(tmp_path):
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    import interface_web
    client = TestClient(interface_web.app)
    pdf = pdf_facture_demo(str(tmp_path / "f.pdf"))
    with open(pdf, "rb") as fh:
        contenu = fh.read()

    def envoyer(**data):
        files = [("fichiers", ("f.pdf", contenu, "application/pdf"))]
        return client.post("/api/extraire", files=files, data=data).json()

    # Usage simple : Excel seul, aucun fichier QuickBooks.
    simple = envoyer()
    assert simple["qbo"] is False
    assert client.get(f"/telecharger/{simple['download_id']}?format=qbo").status_code == 404

    # Option cochée : le CSV s'ajoute, l'Excel reste le même.
    avec = envoyer(qbo="1")
    assert avec["qbo"] is True
    r = client.get(f"/telecharger/{avec['download_id']}?format=qbo")
    assert r.status_code == 200 and "text/csv" in r.headers["content-type"]
    lignes = list(csv.DictReader(io.StringIO(r.text)))
    assert lignes and lignes[0]["Bill no."] == "DEMO-001"

    def contenu_excel(jeton):
        wb = load_workbook(io.BytesIO(client.get("/telecharger/" + jeton).content))
        return {n: [list(r) for r in wb[n].iter_rows(values_only=True)]
                for n in wb.sheetnames}
    assert contenu_excel(simple["download_id"]) == contenu_excel(avec["download_id"])


def test_option_quickbooks_repliee_dans_la_page():
    from fastapi.testclient import TestClient
    import interface_web
    page = TestClient(interface_web.app).get("/").text
    # Options comptables repliées par défaut ; case QuickBooks décochée.
    assert '<details id="avance" class="avance">' in page
    assert '<input id="qboInput" type="checkbox">' in page
    assert "Claude" not in page and "Anthropic" not in page


def test_trousse_de_test_a_jour(tmp_path):
    """Le CSV d'exemple fourni au testeur QuickBooks correspond au code actuel."""
    import importlib.util
    dossier = RACINE / "exemples" / "quickbooks"
    spec = importlib.util.spec_from_file_location("gen", dossier / "generer_exemple_qbo.py")
    gen = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(gen)
    neuf = fve.construire_csv_qbo(gen.factures(), str(tmp_path / "x.csv"), plan=gen.PLAN)
    assert open(neuf, encoding="utf-8").read() == \
        (dossier / "factures_qbo_exemple.csv").read_text(encoding="utf-8")
