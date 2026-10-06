# -*- coding: utf-8 -*-
"""Tests du moteur : extraction, regroupement, ajout à un Excel existant."""
from _helpers import RACINE, pdf_facture_demo  # noqa: F401  (ajuste sys.path)
import facture_vers_excel as fve
from openpyxl import load_workbook


def _fac(nom, four, ttc, n_lignes, date="2026-09-01"):
    d = {"fournisseur": four, "numero": "X", "date": date, "devise": "CAD",
         "total_ht": None, "tps": None, "tvq": None, "tva": None,
         "total_ttc": ttc, "lignes": [{"description": f"l{i}", "montant": 1.0}
                                       for i in range(n_lignes)]}
    return fve._vers_facture(nom, d, "vision:test")


def test_extraction_tables_pdf(tmp_path):
    pdf = pdf_facture_demo(str(tmp_path / "f.pdf"))
    f = fve.traiter_facture(pdf, moteur="tables")
    assert f.moteur == "tables"
    assert f.total_ttc == 200.0
    assert len(f.lignes) >= 2


def test_extraction_llm_simulee(tmp_path, monkeypatch):
    pdf = pdf_facture_demo(str(tmp_path / "f.pdf"))
    faux = {"fournisseur": "Atelier Demo Inc.", "numero": "DEMO-001",
            "date": "2026-09-01", "devise": "CAD", "total_ht": 180.0,
            "tps": None, "tvq": None, "tva": 20.0, "total_ttc": 200.0,
            "lignes": [{"description": "Service A", "quantite": 2,
                        "prix_unitaire": 50, "montant": 100},
                       {"description": "Service B", "quantite": 1,
                        "prix_unitaire": 80, "montant": 80}]}
    monkeypatch.setenv("ANTHROPIC_API_KEY", "cle-test")
    monkeypatch.setattr(fve, "extraire_avec_llm", lambda *a, **k: faux)
    f = fve.traiter_facture(pdf, moteur="llm")
    assert f.total_ttc == 200.0 and len(f.lignes) == 2
    assert not f.alertes  # HT + TVA == TTC, cohérent


def test_regroupement():
    factures = [_fac("recu.heic", "Renaissance Fripe", 151.5, 21),
                _fac("paiement.heic", "Renaissance", 151.5, 0),
                _fac("pharma.heic", "Pharmaprix", 48.42, 4, date="2026-09-02")]
    groupe = fve.regrouper_factures(factures)
    assert len(groupe) == 2
    ren = [f for f in groupe if "Renaissance" in f.fournisseur][0]
    assert ren.pieces_liees == ["paiement.heic"] and len(ren.lignes) == 21


def test_ajout_a_excel_existant(tmp_path):
    base = str(tmp_path / "compta.xlsx")
    fve.construire_excel([_fac("a.pdf", "Alpha", 10, 2)], base)
    fve.construire_excel([_fac("b.jpg", "Beta", 20, 1)], base, base_excel=base)
    wb = load_workbook(base)
    assert wb.sheetnames == ["Resume", "Details"]
    noms = [wb["Resume"].cell(i, 2).value for i in range(2, wb["Resume"].max_row + 1)
            if not fve._est_ligne_total(wb["Resume"].cell(i, 1).value)]
    assert noms == ["Alpha", "Beta"]
    assert wb["Details"].max_row - 1 == 3  # 2 + 1 lignes de détail


def test_ajout_a_feuille_existante_par_titres(tmp_path):
    """Ajout à un onglet Resume aux titres différents : colonnes alignées par
    titre, titres manquants ajoutés, jamais de colonne « Moteur »."""
    from openpyxl import Workbook, load_workbook
    base = str(tmp_path / "ancien.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Resume"
    ws.append(["Fichier", "Fournisseur", "N° facture", "Date", "Total HT", "TVA", "››"])
    ws.append(["a.pdf", "Ancien", "INV-1", "2026-09-15", 100.0, 13.0, 113.0])
    wb.save(base)

    d = {"fournisseur": "Pharmaprix", "numero": "104701", "date": "2026-09-22",
         "devise": "CAD", "total_ht": 43.36, "tps": 1.69, "tvq": 3.37,
         "tva": 5.06, "total_ttc": 48.42,
         "lignes": [{"description": "x", "montant": 43.36}]}
    sortie = str(tmp_path / "sortie.xlsx")
    fve.construire_excel([fve._vers_facture("b.heic", d, "vision:test")],
                         sortie, base_excel=base)

    ws = load_workbook(sortie)["Resume"]
    titres = [c.value for c in ws[1]]
    assert "Moteur" not in titres
    assert all(t for t in titres)                      # aucun titre vide
    ligne = {titres[i]: c.value for i, c in enumerate(ws[3])}
    assert ligne["N° facture"] == "104701"             # retrouvé par synonyme
    assert ligne["Total HT"] == 43.36
    assert ligne["TVA"] == 5.06                        # « Taxes » = « TVA »
    assert ligne["Devise"] == "CAD" and ligne["Total TTC"] == 48.42
    assert "vision:test" not in [c.value for row in ws.iter_rows() for c in row]
    assert ws.cell(2, 2).value == "Ancien"             # ligne existante intacte
    assert ws.cell(4, 1).value.startswith("TOTAL")     # totaux en bas


def _facture_test(nom="c.pdf", numero="42"):
    d = {"fournisseur": "Pharmaprix", "numero": numero, "date": "2026-09-22",
         "devise": "CAD", "total_ht": 43.36, "tps": 1.69, "tvq": 3.37,
         "tva": 5.06, "total_ttc": 99.0,          # incohérent → une alerte
         "lignes": [{"description": "x", "quantite": 1, "montant": 43.36}]}
    return fve._vers_facture(nom, d, "vision:test")


def test_excel_en_anglais(tmp_path):
    from openpyxl import load_workbook
    sortie = str(tmp_path / "en.xlsx")
    fve.construire_excel([_facture_test()], sortie, langue="en")
    wb = load_workbook(sortie)
    assert wb.sheetnames == ["Summary", "Details"]
    titres = [c.value for c in wb["Summary"][1]]
    assert titres[:6] == ["File", "Vendor", "Invoice No.", "Date", "Currency", "Subtotal"]
    assert "Warnings" in titres and "Moteur" not in titres
    alerte = wb["Summary"].cell(2, titres.index("Warnings") + 1).value
    assert alerte.startswith("Subtotal+tax") and "≠ total" in alerte
    assert [c.value for c in wb["Details"][1]][3:] == [
        "Description", "Quantity", "Unit price", "Amount"]


def test_bascule_fr_en_fr(tmp_path):
    """Un Excel français complété en anglais (puis l'inverse) : mêmes colonnes,
    lignes à la suite, le classeur prend la langue du dernier ajout."""
    from openpyxl import load_workbook
    fr = str(tmp_path / "fr.xlsx")
    en = str(tmp_path / "en.xlsx")
    fr2 = str(tmp_path / "fr2.xlsx")
    fve.construire_excel([_facture_test("a.pdf", "1")], fr, langue="fr")
    fve.construire_excel([_facture_test("b.pdf", "2")], en, base_excel=fr, langue="en")

    wb = load_workbook(en)
    assert wb.sheetnames == ["Summary", "Details"]
    ws = wb["Summary"]
    titres = [c.value for c in ws[1]]
    assert titres == [fve._EN[c] for c in fve._COLS_RESUME]   # aucune colonne en double
    assert [r[2] for r in ws.iter_rows(min_row=2, values_only=True)
            if not fve._est_ligne_total(r[0])] == ["1", "2"]
    assert wb["Details"].max_row == 3

    fve.construire_excel([_facture_test("c.pdf", "3")], fr2, base_excel=en, langue="fr")
    wb = load_workbook(fr2)
    assert wb.sheetnames == ["Resume", "Details"]
    assert [c.value for c in wb["Resume"][1]] == fve._COLS_RESUME
    assert wb["Resume"].max_row == 5                  # 3 pièces + 1 TOTAL


def test_totaux_cumules_mis_a_jour(tmp_path):
    """Une ligne TOTAL par devise, recalculée en bas à chaque ajout ;
    « $ » seul compte comme CAD."""
    base = str(tmp_path / "compta.xlsx")
    fve.construire_excel([_facture_test("a.pdf", "1")], base)
    dollar = _facture_test("b.pdf", "2")
    dollar.devise = fve.normaliser_devise("$")
    fve.construire_excel([dollar], base, base_excel=base)
    usd = _facture_test("c.pdf", "3")
    usd.devise = fve.normaliser_devise("US$")
    fve.construire_excel([usd], base, base_excel=base, langue="en")

    ws = load_workbook(base)["Summary"]
    col = [c.value for c in ws[1]]
    fichiers = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    assert fichiers == ["a.pdf", "b.pdf", "c.pdf", "TOTAL CAD", "TOTAL USD"]
    assert ws.cell(5, col.index("Total") + 1).value == 198.0      # valeur, pas formule
    assert ws.cell(5, col.index("Tax") + 1).value == 10.12
    assert ws.cell(5, 2).value == "2 invoice(s)"

    totaux = {t["devise"]: t for t in fve.totaux_resume(base)}
    assert totaux["CAD"]["nb"] == 2 and totaux["CAD"]["total_ttc"] == 198.0
    assert totaux["USD"]["nb"] == 1


def test_normaliser_devise():
    assert [fve.normaliser_devise(v) for v in ("$", "cad", "US$", "€", None)] == \
        ["CAD", "CAD", "USD", "EUR", ""]
