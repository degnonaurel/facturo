# -*- coding: utf-8 -*-
"""Tests de sécurité : injection de formules, limites d'envoi, messages
d'erreur sans détails internes, en-têtes HTTP, expiration des fichiers."""
import csv
from _helpers import RACINE, pdf_facture_demo  # noqa: F401
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import facture_vers_excel as fve
import interface_web

client = TestClient(interface_web.app)
PIEGE = '=HYPERLINK("http://exemple.invalid","Payer ici")'


def _facture_piegee():
    return fve._vers_facture("x.pdf", {
        "fournisseur": PIEGE, "numero": "1", "date": "2026-10-01", "total_ht": 10.0,
        "total_ttc": 10.0, "lignes": [{"description": "@SUM(1+1)", "montant": 10.0}]}, "t")


def test_excel_sans_formule_injectee(tmp_path):
    chemin = fve.construire_excel([_facture_piegee()], str(tmp_path / "f.xlsx"))
    wb = load_workbook(chemin)
    cell = wb["Resume"]["B2"]
    assert cell.value == PIEGE and cell.data_type == "s"      # texte, pas formule
    assert all(c.data_type != "f" for ws in wb for row in ws.iter_rows() for c in row)


def test_csv_qbo_sans_formule_injectee(tmp_path):
    chemin = fve.construire_csv_qbo([_facture_piegee()], str(tmp_path / "q.csv"))
    ligne = next(csv.DictReader(open(chemin, encoding="utf-8")))
    assert ligne["Supplier"] == "'" + PIEGE
    assert ligne["Line Description"] == "'@SUM(1+1)"
    assert fve._texte_csv_sur("-12.50") == "-12.50"           # un nombre reste intact


def test_entetes_de_securite():
    r = client.get("/")
    assert r.headers["X-Content-Type-Options"] == "nosniff"
    assert r.headers["X-Frame-Options"] == "DENY"
    assert "frame-ancestors 'none'" in r.headers["Content-Security-Policy"]


def test_affichage_echappe_les_textes_lus():
    page = client.get("/").text
    assert "function esc(" in page
    assert "${esc(f.fournisseur)" in page and "${esc(e.message)}" in page


def test_fichier_trop_gros_signale_sans_bloquer_le_lot(tmp_path, monkeypatch):
    monkeypatch.setattr(interface_web, "_quota", {"jour": None, "total": 0, "visiteurs": {}})
    pdf = open(pdf_facture_demo(str(tmp_path / "f.pdf")), "rb").read()
    monkeypatch.setattr(interface_web, "TAILLE_MAX_FICHIER", len(pdf) + 10)
    files = [("fichiers", ("ok.pdf", pdf, "application/pdf")),
             ("fichiers", ("gros.pdf", pdf + b"x" * 100, "application/pdf"))]
    d = client.post("/api/extraire", files=files).json()
    assert len(d["factures"]) == 1
    assert d["erreurs"][0]["fichier"] == "gros.pdf"
    assert "volumineux" in d["erreurs"][0]["message"]


def test_trop_de_fichiers(monkeypatch):
    monkeypatch.setattr(interface_web, "MAX_FICHIERS_REQUETE", 2)
    files = [("fichiers", (f"f{i}.pdf", b"x", "application/pdf")) for i in range(3)]
    assert client.post("/api/extraire", files=files).status_code == 413


def test_erreur_de_service_generique(monkeypatch):
    """Le message de l'API (nom du moteur, détails) ne sort pas du serveur."""
    monkeypatch.setattr(interface_web, "_quota", {"jour": None, "total": 0, "visiteurs": {}})
    monkeypatch.setattr(fve, "traiter_lot", lambda chemins, **k: [
        fve.ErreurService("Erreur API Claude 500 : détail interne"),
        ValueError("trace interne")])
    files = [("fichiers", ("a.jpg", b"x", "image/jpeg")),
             ("fichiers", ("b.jpg", b"x", "image/jpeg"))]
    d = client.post("/api/extraire", files=files, data={"langue": "en"}).json()
    textes = " ".join(e["message"] for e in d["erreurs"])
    assert "Claude" not in textes and "interne" not in textes
    assert "temporarily unavailable" in textes and "Unreadable" in textes


def test_fichiers_produits_effaces_apres_expiration(tmp_path, monkeypatch):
    monkeypatch.setattr(interface_web, "_quota", {"jour": None, "total": 0, "visiteurs": {}})
    pdf = open(pdf_facture_demo(str(tmp_path / "f.pdf")), "rb").read()
    d = client.post("/api/extraire", files=[("fichiers", ("f.pdf", pdf, "application/pdf"))],
                    data={"qbo": "1"}).json()
    chemin = interface_web._TELECHARGEMENTS[d["download_id"]][0]
    assert client.get("/telecharger/" + d["download_id"]).status_code == 200
    monkeypatch.setattr(interface_web, "DUREE_TELECHARGEMENT", -1)   # tout est expiré
    assert client.get("/telecharger/" + d["download_id"]).status_code == 404
    import os
    assert not os.path.exists(chemin) and not os.path.exists(chemin[:-5] + "_qbo.csv")
