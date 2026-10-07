# -*- coding: utf-8 -*-
"""Tests des relevés bancaires (CSV/OFX) et du rapprochement avec les factures."""
from _helpers import RACINE  # noqa: F401
from openpyxl import load_workbook
import facture_vers_excel as fve
import releves_bancaires as rb


def _ecrire(tmp_path, nom, texte, codage="utf-8"):
    p = tmp_path / nom
    p.write_bytes(texte.encode(codage))
    return str(p)


# --- Lecture des formats -------------------------------------------------

def test_csv_avec_titres_montant_signe(tmp_path):
    """Style RBC : titres anglais, montant signé, deux colonnes de description."""
    p = _ecrire(tmp_path, "rbc.csv",
                '"Account Type","Account Number","Transaction Date","Cheque Number",'
                '"Description 1","Description 2","CAD$","USD$"\n'
                'Chequing,12345,9/3/2026,,"PAPETERIE GATIN","GATINEAU QC",-137.97,\n'
                'Chequing,12345,9/15/2026,,"DEPOT PAIE",,2500.00,\n')
    t = rb.lire_releve(p)
    assert [(x.date, x.montant) for x in t] == [("2026-09-03", -137.97),
                                               ("2026-09-15", 2500.0)]
    assert t[0].description == "PAPETERIE GATIN GATINEAU QC"


def test_csv_sans_titres_debit_credit_solde(tmp_path):
    """Style TD : pas de titres, MM/JJ/AAAA, colonnes débit / crédit / solde."""
    p = _ecrire(tmp_path, "td.csv",
                "09/05/2026,BISTRO DU RUISS,73.58,,1926.42\n"
                "09/14/2026,VIREMENT RECU,,500.00,2426.42\n")
    t = rb.lire_releve(p)
    assert [(x.date, x.montant) for x in t] == [("2026-09-05", -73.58),
                                               ("2026-09-14", 500.0)]


def test_csv_francais_point_virgule_virgule_decimale(tmp_path):
    """Style Desjardins/Banque Nationale : français, « ; », virgule décimale,
    JJ/MM/AAAA, colonnes Retrait / Dépôt, codage Windows."""
    p = _ecrire(tmp_path, "desj.csv",
                "Date;Description;Retrait;Dépôt;Solde\n"
                "28/09/2026;Rideau Hardware Ottawa;282,50;;1 717,50\n"
                "30/09/2026;Dépôt mobile;;1 000,00;2 717,50\n", codage="cp1252")
    t = rb.lire_releve(p)
    assert [(x.date, x.montant) for x in t] == [("2026-09-28", -282.5),
                                               ("2026-09-30", 1000.0)]
    assert t[1].description == "Dépôt mobile"


def test_ofx_sgml(tmp_path):
    p = _ecrire(tmp_path, "releve.ofx", """OFXHEADER:100
DATA:OFXSGML
<OFX><BANKMSGSRSV1><STMTTRNRS><STMTRS><CURDEF>CAD
<BANKTRANLIST>
<STMTTRN><TRNTYPE>DEBIT<DTPOSTED>20260910120000[-5:EST]<TRNAMT>-315.00
<FITID>A1<NAME>FORMATION PRO ALB<MEMO>Achat
<STMTTRN><TRNTYPE>CREDIT<DTPOSTED>20260911<TRNAMT>40.00<FITID>A2<NAME>REMBOURSEMENT
</BANKTRANLIST></STMTRS></STMTTRNRS></BANKMSGSRSV1></OFX>""")
    t = rb.lire_releve(p)
    assert [(x.date, x.montant, x.devise) for x in t] == [
        ("2026-09-10", -315.0, "CAD"), ("2026-09-11", 40.0, "CAD")]
    assert t[0].description == "FORMATION PRO ALB Achat" and t[0].reference == "A1"


def test_ordre_des_dates_et_montants():
    assert rb.ordre_des_dates(["03/09/2026", "25/09/2026"]) == "jm"
    assert rb.ordre_des_dates(["09/03/2026", "09/25/2026"]) == "mj"
    assert rb.lire_montant("(12.50)") == -12.5
    assert rb.lire_montant("1 234,56 $") == 1234.56


def test_releve_illisible(tmp_path):
    import pytest
    p = _ecrire(tmp_path, "vide.csv", "rien,ici\nni,montant\n")
    with pytest.raises(rb.ErreurReleve):
        rb.lire_releve(p)


# --- Rapprochement ---------------------------------------------------------

def _piece(fournisseur, date, total, numero="1"):
    return {"fichier": f"{numero}.pdf", "fournisseur": fournisseur, "numero": numero,
            "date": date, "devise": "CAD", "total_ttc": total, "categorie": ""}


def test_rapprochement_montant_date_et_nom():
    tr = [rb.Transaction("2026-09-04", "PAPETERIE GATIN", -137.97, "CAD"),
          rb.Transaction("2026-09-05", "STATION ESSO", -60.00, "CAD"),     # sans pièce
          rb.Transaction("2026-10-20", "RIDEAU HARDWARE", -282.50, "CAD"),  # facture payée à 42 j
          rb.Transaction("2026-09-06", "DEPOT", 137.97, "CAD")]             # entrée : ignorée
    pieces = [_piece("Papeterie Gatineau", "2026-09-03", 137.97, "PG"),
              _piece("Rideau Hardware", "2026-09-08", 282.50, "RH"),
              _piece("Bistro du Ruisseau", "2026-09-05", 73.58, "BR")]       # absente
    r = rb.rapprocher(tr, pieces)
    assert [(i, k) for i, k, *_ in r["paires"]] == [(0, 0), (2, 1)]
    assert r["paires"][0][3] == "elevee"
    assert r["transactions_seules"] == [1, 3] and r["pieces_seules"] == [2]


def test_meme_montant_sans_nom_trop_loin_refuse():
    """Un montant identique, sans nom reconnu et 30 jours plus tard : pas un match."""
    tr = [rb.Transaction("2026-10-03", "VIREMENT INTERAC", -50.0)]
    r = rb.rapprocher(tr, [_piece("Fleuriste Rose", "2026-09-03", 50.0)])
    assert r["paires"] == []


def test_deux_achats_meme_montant_departages_par_le_nom():
    tr = [rb.Transaction("2026-09-10", "TIM HORTONS", -12.50),
          rb.Transaction("2026-09-10", "PAPETERIE GATIN", -12.50)]
    pieces = [_piece("Papeterie Gatineau", "2026-09-10", 12.50, "P"),
              _piece("Tim Hortons", "2026-09-10", 12.50, "T")]
    r = rb.rapprocher(tr, pieces)
    assert sorted((i, k) for i, k, *_ in r["paires"]) == [(0, 1), (1, 0)]


# --- Onglet Excel ------------------------------------------------------------

def test_onglet_rapprochement_sur_tout_le_classeur(tmp_path):
    """Le relevé est rapproché de TOUTES les pièces du classeur, y compris
    celles ajoutées lors d'envois précédents."""
    def facture(fournisseur, numero, date, total):
        return fve._vers_facture(f"{numero}.pdf", {
            "fournisseur": fournisseur, "numero": numero, "date": date,
            "devise": "CAD", "total_ttc": total, "lignes": []}, "t")
    base = str(tmp_path / "compta.xlsx")
    fve.construire_excel([facture("Papeterie Gatineau", "PG", "2026-09-03", 137.97)], base,
                         plan=fve.plan_par_defaut("fr"))

    tr = [rb.Transaction("2026-09-04", "PAPETERIE GATIN", -137.97, "CAD"),
          rb.Transaction("2026-09-05", '=HYPERLINK("x")', -60.00, "CAD")]
    bilan = {}
    sortie = str(tmp_path / "compta2.xlsx")
    fve.construire_excel([facture("Bistro du Ruisseau", "BR", "2026-09-05", 73.58)], sortie,
                         base_excel=base, releve=tr, bilan=bilan)
    wb = load_workbook(sortie)
    assert wb.sheetnames == ["Resume", "Details", "Totaux par catégorie",
                             "Rapprochement", "Plan comptable"]
    lignes = list(wb["Rapprochement"].iter_rows(min_row=2, values_only=True))
    statuts = [l[4] for l in lignes]
    assert statuts == ["Rapprochée", "Sans pièce", "Absente du relevé"]
    assert lignes[0][5] == "PG.pdf" and lignes[2][6] == "Bistro du Ruisseau"
    assert bilan == {"transactions": 2, "rapprochees": 1, "sans_piece": 1,
                     "montant_sans_piece": 60.0, "categorisees": 0, "absentes": 1}
    assert wb["Rapprochement"]["B3"].data_type == "s"     # libellé bancaire = texte


def test_onglet_en_anglais(tmp_path):
    f = fve._vers_facture("a.pdf", {"fournisseur": "A", "numero": "1",
                                    "date": "2026-09-01", "total_ttc": 5.0, "lignes": []}, "t")
    sortie = str(tmp_path / "en.xlsx")
    fve.construire_excel([f], sortie, langue="en",
                         releve=[rb.Transaction("2026-09-01", "A", -5.0)])
    ws = load_workbook(sortie)["Bank reconciliation"]
    assert ws["E1"].value == "Status" and ws["E2"].value == "Matched"


# --- App web -----------------------------------------------------------------

def _client(monkeypatch):
    from fastapi.testclient import TestClient
    import interface_web
    monkeypatch.setattr(interface_web, "_quota", {"jour": None, "total": 0, "visiteurs": {}})
    return TestClient(interface_web.app)


RELEVE_CSV = (b"Date,Description,Amount\n"
              b"2026-09-01,ATELIER DEMO INC,-200.00\n"
              b"2026-09-02,STATION ESSO,-45.00\n")


def test_web_facture_et_releve(tmp_path, monkeypatch):
    from _helpers import pdf_facture_demo
    client = _client(monkeypatch)
    pdf = open(pdf_facture_demo(str(tmp_path / "f.pdf")), "rb").read()
    d = client.post("/api/extraire", files=[
        ("fichiers", ("f.pdf", pdf, "application/pdf")),
        ("releve", ("releve.csv", RELEVE_CSV, "text/csv"))]).json()
    assert d["rapprochement"]["rapprochees"] == 1
    assert d["rapprochement"]["sans_piece"] == 1
    assert d["rapprochement"]["montant_sans_piece"] == 45.0


def test_web_releve_seul_avec_excel_existant(tmp_path, monkeypatch):
    """Le comptable ajoute seulement le relevé du mois à son classeur."""
    client = _client(monkeypatch)
    base = str(tmp_path / "base.xlsx")
    f = fve._vers_facture("a.pdf", {"fournisseur": "Atelier Demo Inc.", "numero": "1",
                                    "date": "2026-09-01", "devise": "CAD",
                                    "total_ttc": 200.0, "lignes": []}, "t")
    fve.construire_excel([f], base)
    import interface_web
    avant = interface_web._quota["total"]
    d = client.post("/api/extraire", files=[
        ("base", ("base.xlsx", open(base, "rb").read(), "application/octet-stream")),
        ("releve", ("releve.csv", RELEVE_CSV, "text/csv"))]).json()
    assert d["download_id"] and d["factures"] == []
    assert d["rapprochement"]["rapprochees"] == 1
    assert interface_web._quota["total"] == avant        # aucun fichier lu : quota intact


def test_web_releve_illisible_et_envoi_vide(monkeypatch):
    client = _client(monkeypatch)
    d = client.post("/api/extraire", files=[
        ("releve", ("releve.csv", b"rien a voir\n", "text/csv"))]).json()
    assert d["download_id"] is None and "illisible" in d["erreurs"][0]["message"]
    assert client.post("/api/extraire", data={"langue": "fr"}).status_code == 400



# --- Limites levées : devises, catégories, adresse du visiteur ----------------

def test_facture_usd_payee_en_cad():
    tr = [rb.Transaction("2026-09-13", "CLOUDNOTE INC SAN FRANCISCO", -40.12),
          rb.Transaction("2026-09-13", "AUTRE ACHAT", -40.12)]
    piece = dict(_piece("CloudNote Inc.", "2026-09-12", 29.00), devise="USD")
    r = rb.rapprocher(tr, [piece])
    assert r["paires"] == [(0, 0, 1, "converti")]          # le nom confirme


def test_conversion_refusee_sans_nom_ou_hors_fourchette():
    piece = dict(_piece("CloudNote Inc.", "2026-09-12", 29.00), devise="USD")
    assert rb.rapprocher([rb.Transaction("2026-09-13", "ACHAT WEB", -40.12)],
                         [piece])["paires"] == []          # pas de nom
    assert rb.rapprocher([rb.Transaction("2026-09-13", "CLOUDNOTE", -80.00)],
                         [piece])["paires"] == []          # taux 2,76 : impossible


def test_categorie_suggeree_plan_defaut_et_plan_importe():
    defaut = fve.plan_par_defaut("fr")
    assert rb.categorie_suggeree("STATION ESSO OTTAWA", defaut) == "5310 · Carburant"
    assert rb.categorie_suggeree("TIM HORTONS #2231", defaut) == "5210 · Repas et représentation"
    assert rb.categorie_suggeree("FRAIS MENSUELS COMPTE", defaut).startswith("5230")
    assert rb.categorie_suggeree("VIREMENT 4471", defaut) == ""
    qbo = [{"code": "", "nom": "Meals and entertainment", "type": "Expense"},
           {"code": "", "nom": "Bank charges", "type": "Expense"}]
    assert rb.categorie_suggeree("Starbucks 0042", qbo) == "Meals and entertainment"
    assert rb.categorie_suggeree("Esso", qbo) == ""         # pas de compte carburant


def test_onglet_categorie_suggeree(tmp_path):
    f = fve._vers_facture("a.pdf", {"fournisseur": "A", "numero": "1", "date": "2026-09-01",
                                    "total_ttc": 5.0, "lignes": []}, "t")
    bilan, sortie = {}, str(tmp_path / "s.xlsx")
    fve.construire_excel([f], sortie, plan=fve.plan_par_defaut("fr"), bilan=bilan,
                         releve=[rb.Transaction("2026-09-02", "PETRO-CANADA 123", -50.0)])
    ligne = list(load_workbook(sortie)["Rapprochement"].iter_rows(min_row=2, values_only=True))[0]
    assert ligne[4] == "Sans pièce" and ligne[10] == "5310 · Carburant"
    assert ligne[11] == "Catégorie suggérée" and bilan["categorisees"] == 1


def test_visiteur_lu_dans_l_entete_cloudflare():
    from starlette.requests import Request
    import interface_web

    def req(**entetes):
        return Request({"type": "http", "client": ("10.0.0.1", 1), "headers": [
            (k.replace("_", "-").encode(), v.encode()) for k, v in entetes.items()]})
    # Cloudflare écrase CF-Connecting-IP : une fausse valeur dans
    # X-Forwarded-For ne change plus l'identité du visiteur.
    assert interface_web._visiteur(req(cf_connecting_ip="1.2.3.4",
                                       x_forwarded_for="9.9.9.9")) == "1.2.3.4"
    assert interface_web._visiteur(req(x_forwarded_for="5.6.7.8, 10.0.0.2")) == "5.6.7.8"
    assert interface_web._visiteur(req()) == "10.0.0.1"
